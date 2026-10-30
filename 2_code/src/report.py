"""Step 7: turn the results into things people can read.

Builds:
  4_outputs/RESULTS.xlsx     formatted workbook, 'Start here' sheet first
  1_report/report.html       parts 1-6 with charts (print to PDF from any browser)
  1_report/charts/*.png      every chart as an image
  5_finals/slides.html       the six finals slides (print to PDF, landscape)
"""
import base64
import io
import os

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
import pandas as pd  # noqa: E402

INK, MUTED, ACCENT, GOOD, BAD, GRID = "#1f2933", "#7b8794", "#0b6e99", "#2f855a", "#c53030", "#e4e7eb"
plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 10, "axes.edgecolor": GRID,
                     "axes.labelcolor": INK, "xtick.color": INK, "ytick.color": INK,
                     "axes.spines.top": False, "axes.spines.right": False})


# ---------- formatting helpers ----------

def twd(x):
    x = float(x)
    if abs(x) >= 1e6:
        return f"NT${x / 1e6:,.1f}M"
    if abs(x) >= 1e5:
        return f"NT${x / 1e3:,.0f}K"
    return f"NT${x:,.0f}"


def pct(x, d=1):
    return f"{float(x) * 100:.{d}f}%"


def pp(x, d=1):
    return f"{float(x) * 100:.{d}f}"


def n(x):
    return f"{int(round(float(x))):,}"


# ---------- charts ----------

class Charts:
    def __init__(self, folder):
        self.folder = folder
        os.makedirs(folder, exist_ok=True)
        self.b64 = {}

    def save(self, fig, name):
        fig.tight_layout()
        fig.savefig(os.path.join(self.folder, name + ".png"), dpi=160)
        buf = io.BytesIO()
        fig.savefig(buf, format="png", dpi=160)
        plt.close(fig)
        self.b64[name] = base64.b64encode(buf.getvalue()).decode()

    def img(self, name, width="100%"):
        return f'<img src="data:image/png;base64,{self.b64[name]}" style="width:{width}">'


def bar(ax, labels, values, colors, fmt):
    bars = ax.barh(labels, values, color=colors)
    for b, v in zip(bars, values):
        ax.text(b.get_width(), b.get_y() + b.get_height() / 2, " " + fmt(v), va="center", fontsize=9, color=INK)
    ax.invert_yaxis()
    ax.xaxis.set_visible(False)
    ax.spines["bottom"].set_visible(False)


def make_charts(R, ch):
    bt = R["bt_pooled"]
    order = ["R1_leaving_now", "R2_wavering", "R3_new_and_fragile", "R4_unhappy_quiet",
             "BASELINE_contract_ends_<=4m", "ALL_postpaid", "R0_no_signal"]
    names = {"R1_leaving_now": "R1 Note: leaving now", "R2_wavering": "R2 Note: considering leaving",
             "R3_new_and_fragile": "R3 Joined < 6 months ago", "R4_unhappy_quiet": "R4 Complaint, no leave talk",
             "BASELINE_contract_ends_<=4m": "Today's rule: contract ends ≤ 4 months",
             "ALL_postpaid": "All postpaid customers", "R0_no_signal": "No signal"}
    b = bt.set_index("rule").loc[order]
    fig, ax = plt.subplots(figsize=(7.5, 3.2))
    cols = [ACCENT, ACCENT, "#5fa8c9", "#9cc9dd", BAD, MUTED, GRID]
    bar(ax, [names[r] for r in b.index], b["churn_rate"], cols, lambda v: pct(v))
    ax.set_title("Left within 60 days, by rule (3 past test dates)", loc="left", fontsize=10)
    ch.save(fig, "1_backtest")

    m = R["monthly_churn"]
    fig, ax = plt.subplots(figsize=(7.5, 2.6))
    ax.plot(m["month"], m["churn_rate"] * 100, marker="o", color=ACCENT)
    for x, y in zip(m["month"], m["churn_rate"] * 100):
        ax.text(x, y + 0.05, f"{y:.2f}%", ha="center", fontsize=8)
    ax.set_ylim(0, max(m["churn_rate"] * 100) * 1.3)
    ax.set_ylabel("% of postpaid who left")
    ax.grid(axis="y", color=GRID)
    ax.set_title("Postpaid churn per month, 2026", loc="left", fontsize=10)
    ch.save(fig, "2_monthly_churn")

    rs = R["reasons"]
    fig, ax = plt.subplots(figsize=(7.5, 2.6))
    bar(ax, rs["reason"], rs["customers"], ACCENT, lambda v: n(v))
    ax.set_title("Reasons given by flagged customers (R1 + R2)", loc="left", fontsize=10)
    ch.save(fig, "3_reasons")

    roas = R["roas"]
    fig, ax = plt.subplots(figsize=(7.5, 2.8))
    x = range(len(roas))
    ax.bar([i - 0.2 for i in x], roas["reported_roas"], width=0.4, color=MUTED, label="Reported ROAS")
    ax.bar([i + 0.2 for i in x], roas["incremental_roas"], width=0.4, color=ACCENT, label="Real ROAS (vs control group)")
    for i, (a, c) in enumerate(zip(roas["reported_roas"], roas["incremental_roas"])):
        ax.text(i - 0.2, a + 0.05, f"{a:.1f}×", ha="center", fontsize=9)
        ax.text(i + 0.2, c + 0.05, f"{c:.1f}×", ha="center", fontsize=9)
    ax.axhline(1 / 0.6, color=BAD, ls="--", lw=1, label="Break-even on 60% margin (1.67×)")
    ax.set_ylim(0, max(roas["reported_roas"]) * 1.25)
    ax.set_xticks(list(x))
    ax.set_xticklabels(roas["objective"].str.capitalize())
    ax.legend(frameon=False, fontsize=8)
    ax.set_title("Reported vs real ROAS (owned campaigns with a control group)", loc="left", fontsize=10)
    ch.save(fig, "4_roas")

    p = R["paid_by_ch"].sort_values("cost_per_line")
    fig, ax = plt.subplots(figsize=(7.5, 2.4))
    bar(ax, p["channel"], p["cost_per_line"], [GOOD] + [ACCENT] * (len(p) - 2) + [BAD], lambda v: f"NT${v:,.0f}")
    ax.set_title("Cost per new line, paid channels (24 months)", loc="left", fontsize=10)
    ch.save(fig, "5_paid")

    pl = R["pilot"]
    fig, ax = plt.subplots(figsize=(7.5, 2.8))
    x = range(len(pl))
    ax.bar([i - 0.2 for i in x], pl["total_cost"] / 1e6, width=0.4, color=MUTED, label="Pilot cost")
    ax.bar([i + 0.2 for i in x], pl["credited_revenue"] * 0.6 / 1e6, width=0.4, color=GOOD, label="Margin on kept customers")
    for i, (a, c) in enumerate(zip(pl["total_cost"], pl["credited_revenue"] * 0.6)):
        ax.text(i - 0.2, a / 1e6 + 0.05, twd(a), ha="center", fontsize=8)
        ax.text(i + 0.2, c / 1e6 + 0.05, twd(c), ha="center", fontsize=8)
    ax.axhline(2.5, color=BAD, ls="--", lw=1, label="Budget NT$2.5M")
    ax.set_ylim(0, max(pl["credited_revenue"] * 0.6 / 1e6) * 1.2)
    ax.legend(frameon=False, fontsize=8, loc="upper left")
    ax.set_xticks(list(x))
    ax.set_xticklabels([f"{s} case\n(keep {pct(r, 0)})" for s, r in zip(pl["scenario"], pl["save_rate"])])
    ax.set_ylabel("NT$ million")
    ax.set_title("90-day pilot: cost vs value", loc="left", fontsize=10)
    ch.save(fig, "6_pilot")


# ---------- workbook ----------

FRIENDLY = {
    "rule": "Rule", "customers": "Customers", "left_next_60d": "Left within 60 days", "churn_rate": "Churn rate (60 days)",
    "lift_vs_all": "× vs all postpaid", "share_of_all_leavers": "Share of all leavers caught", "cutoff": "Tested on",
    "audience": "Audience", "offer": "Offer", "customers_in_sample": "Customers (sample)",
    "customers_full_base": "Customers (full base, ×130)", "campaign": "Campaign", "who": "Who", "channel": "Channel",
    "channel_plan": "Channel plan", "customers_at_full_scale": "Customers (full base, ×130)", "timing": "Timing",
    "headline": "Headline", "body": "Message", "cta": "Call to action", "compliant": "Compliant", "problems": "Problems",
    "objective": "Objective", "scenario": "Scenario", "save_rate": "Save rate", "take_rate": "Take rate",
    "credited_revenue_twd": "Credited revenue NT$", "offer_cost_twd": "Offer cost NT$", "net_margin_twd": "Net margin NT$",
    "file": "File", "step": "Step", "rows_before": "Rows before", "rows_after": "Rows after", "changed": "Rows removed",
    "note": "Note", "campaign_id": "Campaign", "month": "Month", "target_segment": "Target as written",
    "vague_target": "Vague target", "spend": "Spend NT$", "reported_revenue": "Reported revenue NT$",
    "reported_roas": "Reported ROAS", "incremental_roas": "Real ROAS (vs control)",
    "incremental_margin_roas": "Real margin return", "share_of_reported_that_is_real": "Share of reported that is real",
    "cost_per_line": "Cost per new line NT$", "new_lines": "New lines", "margin_roas_6m": "6-month margin return",
    "median_payback_months": "Payback months (median)", "plan_type": "Plan", "tier": "Tier",
    "tier_reason": "Why", "avg_bill": "Average bill NT$", "key": "Key", "what": "What", "value": "Value", "why": "Why",
}

SHEETS = [
    ("Backtest", "bt_pooled"), ("Audiences", "aud"), ("New campaigns", "camp"), ("Pilot", "pilot"),
    ("Value ranges", "val"), ("Past campaigns real ROAS", "roas"), ("Paid channels", "paid_by_ch"),
    ("Copy compliance audit", "copy_audit"), ("Segments", "seg"), ("Cleaning log", "log"), ("Assumptions", "assump"),
]


def write_workbook(R, path, headline_rows):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        pd.DataFrame(headline_rows, columns=["What we found", "Number"]).to_excel(xw, sheet_name="Start here", index=False)
        for sheet, key in SHEETS:
            df = R[key].copy()
            df = df.rename(columns={c: FRIENDLY.get(c, c.replace("_", " ").capitalize()) for c in df.columns})
            df.to_excel(xw, sheet_name=sheet, index=False)
        for ws in xw.book.worksheets:
            ws.freeze_panes = "A2"
            for c in ws[1]:
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = PatternFill("solid", fgColor="0B6E99")
                c.alignment = Alignment(wrap_text=True, vertical="top")
            for i, col in enumerate(ws.columns, 1):
                head = str(col[0].value or "")
                width = min(max(len(head), *(len(str(c.value or "")) for c in col[1:60])) + 2, 60)
                ws.column_dimensions[get_column_letter(i)].width = max(width, 10)
                for c in col[1:]:
                    if isinstance(c.value, float):
                        h = head.lower()
                        if "rate" in h or "share" in h:
                            c.number_format = "0.0%"
                        elif "nt$" in h or "revenue" in h or "cost" in h or "spend" in h:
                            c.number_format = "#,##0"
                        else:
                            c.number_format = "#,##0.00"
                    elif isinstance(c.value, int):
                        c.number_format = "#,##0"
                    if isinstance(c.value, str) and len(c.value) > 60:
                        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws = xw.book["Start here"]
        ws.column_dimensions["A"].width = 90
        ws.column_dimensions["B"].width = 28


# ---------- HTML ----------

CSS = """
@page { size: A4; margin: 16mm 15mm; }
body { font-family: 'Segoe UI', Helvetica, Arial, sans-serif; color: #1f2933; font-size: 10.5pt; line-height: 1.45; max-width: 820px; margin: auto; }
h1 { font-size: 22pt; margin: 0 0 4px; } h2 { font-size: 15pt; color: #0b6e99; border-bottom: 2px solid #0b6e99; padding-bottom: 3px; margin-top: 0; }
h3 { font-size: 11.5pt; margin: 14px 0 4px; page-break-after: avoid; } .sub { color: #7b8794; margin-bottom: 14px; }
.part { page-break-before: always; } .first { page-break-before: auto; }
table { border-collapse: collapse; width: 100%; margin: 6px 0 12px; font-size: 9pt; } tr { page-break-inside: avoid; } img { page-break-inside: avoid; }
th { background: #0b6e99; color: white; text-align: left; padding: 4px 6px; } td { border-bottom: 1px solid #e4e7eb; padding: 4px 6px; vertical-align: top; }
td.num, th.num { text-align: right; } .box { background: #f0f7fa; border-left: 4px solid #0b6e99; padding: 8px 12px; margin: 8px 0 12px; }
.ask { background: #fff8e6; border-left: 4px solid #d69e2e; padding: 8px 12px; } .kpis { display: flex; gap: 8px; margin: 8px 0 12px; }
.kpi { flex: 1; border: 1px solid #e4e7eb; border-radius: 6px; padding: 8px; } .kpi b { font-size: 14pt; color: #0b6e99; display: block; }
.small { font-size: 8.5pt; color: #52606d; } img { margin: 4px 0 8px; } ul { margin: 4px 0 8px 18px; padding: 0; }
"""


def table(df, cols, heads=None, fmts=None):
    heads = heads or cols
    fmts = fmts or {}
    th = "".join(f'<th class="{"num" if c in fmts else ""}">{h}</th>' for c, h in zip(cols, heads))
    rows = []
    for _, r in df.iterrows():
        tds = "".join(f'<td class="{"num" if c in fmts else ""}">{fmts[c](r[c]) if c in fmts and pd.notna(r[c]) else ("" if pd.isna(r[c]) else r[c])}</td>' for c in cols)
        rows.append(f"<tr>{tds}</tr>")
    return f"<table><tr>{th}</tr>{''.join(rows)}</table>"


def build(R, root):
    ch = Charts(os.path.join(root, "1_report", "charts"))
    make_charts(R, ch)
    K = R["K"]
    hc = os.path.join(root, "3_ai_outputs", "hand_check", "result.json")
    if os.path.exists(hc):
        import json
        with open(hc, encoding="utf-8") as f:
            h = json.load(f)
        K["hand_check_text"] = (f"{h['labelers']} of us labelled {h['sentences']} sentences blind; we agreed with the AI on "
                                f"leave / not leave for {h['flag_agree'] * 100:.0f}% of them and on the topic for {h['topic_agree'] * 100:.0f}%.")
    else:
        K["hand_check_text"] = ("Two of us are labelling a blind sample of 100 of those sentences to measure agreement; "
                                "see 3_ai_outputs/hand_check.")

    headline = [
        ("Customers whose support notes say they may leave: churn in next 60 days", pct(K["sig_rate"])),
        ("All postpaid customers: churn in next 60 days", pct(K["base_rate"])),
        ("Formosa's current rule (contract ends ≤ 4 months): churn in next 60 days", pct(K["ce_rate"])),
        ("Notes that say 'may leave' but were not tagged cancel_inquiry by the agent", n(K["missed"])),
        ("Objective 1, churn: net margin per year at full scale (low – high)", f"{twd(K['churn_low'])} – {twd(K['churn_high'])}"),
        ("Objective 2, earn more: net margin per year incl. roaming spend saved (low – high)", f"{twd(K['up_low'])} – {twd(K['up_high'])}"),
        ("Objective 3, win new: extra 6-month margin per year from channel shift (low – high spend)", f"{twd(K['acq_low'])} – {twd(K['acq_high'])}"),
        ("Pilot cost, 90 days (low – high case)", f"{twd(K['pilot_cost_low'])} – {twd(K['pilot_cost_high'])}"),
        ("Pilot break-even: share of would-be leavers we must keep", pct(K["breakeven"])),
        ("Past campaigns that break brand or legal rules", f"{K['bad_copy']} of {K['all_copy']}"),
        ("Roaming campaigns: extra conversions vs control group", "none measurable"),
    ]
    write_workbook(R, os.path.join(root, "4_outputs", "RESULTS.xlsx"), headline)

    html = report_html(R, K, ch)
    with open(os.path.join(root, "1_report", "report.html"), "w", encoding="utf-8") as f:
        f.write(html)
    with open(os.path.join(root, "5_finals", "slides.html"), "w", encoding="utf-8") as f:
        f.write(slides_html(R, K, ch))
    with open(os.path.join(root, "5_finals", "talk_script.md"), "w", encoding="utf-8") as f:
        f.write(talk_script(R, K))
    with open(os.path.join(root, "4_outputs", "campaign_builder.html"), "w", encoding="utf-8") as f:
        f.write(campaign_builder(R, os.path.join(root, "3_ai_outputs", "copy_templates.json")))


RULE_NAMES = {"R1_leaving_now": "R1 · note: leaving now", "R2_wavering": "R2 · note: considering leaving",
              "R3_new_and_fragile": "R3 · joined < 6 months ago", "R4_unhappy_quiet": "R4 · complaint, no leave talk",
              "R0_no_signal": "No signal", "X_no_save": "Excluded: line closing, no save",
              "BASELINE_contract_ends_<=4m": "Today's rule: contract ends ≤ 4 months", "ALL_postpaid": "All postpaid customers"}
PLAN_NAMES = {"prepaid": "Prepaid", "postpaid_basic": "Postpaid basic", "postpaid_premium": "Postpaid premium",
              "family": "Family (other team)", "business": "Business (other team)"}
RULE_ORDER = list(RULE_NAMES)


def report_html(R, K, ch):
    bt = R["bt_pooled"].copy()
    bt["order"] = bt["rule"].map({r: i for i, r in enumerate(RULE_ORDER)})
    bt = bt.sort_values("order")
    bt["rule"] = bt["rule"].map(RULE_NAMES)
    R = dict(R, bt_table=bt)
    log = R["log"]
    log = log[(log["changed"] > 0) | log["note"].astype(str).ne("")]
    camp = R["camp"].copy()
    camp["compliant"] = camp["compliant"].map({True: "Yes", False: "No"})
    ex = camp[camp["offer"] == "Renewal discount"].iloc[0] if (camp["offer"] == "Renewal discount").any() else camp.iloc[0]
    roas = R["roas"]
    p = R["pilot"]
    mid = p[p["scenario"] == "mid"].iloc[0]
    aud = R["aud"]
    seg = R["seg_plan"]
    parts = []

    # ---------------- Part 1 ----------------
    parts.append(f"""
<div class="first">
<h1>Formosa Telecom: listen to what customers tell support</h1>
<div class="sub">DeployAI 2026 · Commercial proposal for Vivian Hsu, CMO · data to 30 June 2026</div>
<h2>1. Executive summary</h2>
<div class="box"><b>The finding that changes the plan.</b> When a customer tells a support agent they are thinking of leaving,
about half of them are gone within 60 days ({pct(K['sig_rate'])} vs {pct(K['base_rate'])} for all postpaid customers, tested on three past dates
using only notes written before each date). Formosa's current renewal push, aimed at contracts ending within four months,
reaches people who leave at {pct(K['ce_rate'])}: barely above average. Nobody reads these notes today, and {n(K['missed'])} of them were not even tagged as a cancellation risk.</div>
<div class="kpis">
<div class="kpi"><b>{twd(K['churn_low'])}–{twd(K['churn_high'])}</b>a year, keep customers who say they may leave</div>
<div class="kpi"><b>{twd(K['up_low'])}–{twd(K['up_high'])}</b>a year, sell only what customers ask for</div>
<div class="kpi"><b>{twd(K['acq_low'])}–{twd(K['acq_high'])}</b>a year, move paid spend to the cheapest channel</div>
</div>
<p class="small">Net margin at Formosa's full scale (sample ×130), 60% margin, six months of credited revenue per customer as Finance counts it. Ranges and assumptions in Part 4.</p>
<h3>1 · Reduce churn</h3>
<p><b>Do:</b> read every support note each week, flag postpaid customers who say they may leave, and answer their reason with the approved offer that fits it
(Renewal discount, Data boost or Device credit), calling the highest-bill customers first. Billing errors and network problems get fixed, not discounted.
<b>Worth:</b> {twd(K['churn_low'])} to {twd(K['churn_high'])} net margin a year, keeping {n(K['kept_low'])}–{n(K['kept_high'])} customers.</p>
<h3>2 · Earn more from existing customers</h3>
<p><b>Do:</b> stop roaming blasts: against their control group they caused no extra sales ({twd(K['roam_spend_year'])} a year). Broad upsell blasts report {roas.loc[roas.objective=='upsell','reported_roas'].iloc[0]:.1f}× but really return
{roas.loc[roas.objective=='upsell','incremental_roas'].iloc[0]:.1f}×, below break-even on margin. Send upgrade, data add-on and roaming offers only to customers who asked for them in a note.
<b>Worth:</b> {twd(K['up_low'])} to {twd(K['up_high'])} a year, mostly from spend no longer wasted.</p>
<h3>3 · Win new customers</h3>
<p><b>Do:</b> shift paid budget toward Google Search (cheapest line, {twd(K['best_cpl'])}) and away from Retail Promotion ({twd(K['worst_cpl'])}); every paid channel takes about six months
to pay back, and 2026 campaigns pay back slower than 2025. Add a welcome check-in: customers in their first six months leave at {pct(K['r3_rate'])} in 60 days, {K['r3_lift']:.1f}× the average.
<b>Worth:</b> {twd(K['acq_low'])} to {twd(K['acq_high'])} more margin a year at NT$20–60M of paid spend.</p>
<div class="ask"><b>The one decision:</b> approve a 90-day retention pilot for postpaid customers who tell support they may leave.
Expected cost {twd(mid['total_cost'])} (at most {twd(K['pilot_cost_high'])}, inside the NT$2.5M budget). One in five is held back at random to prove the result.
It pays for itself if it keeps {pct(K['breakeven'])} of would-be leavers, and can prove a save rate of {pct(K['mde_save'])} or more. Paid offers stop at day 45 if contacted customers leave as often as held-back ones.</div>
</div>""")

    # ---------------- Part 2 ----------------
    parts.append(f"""
<div class="part"><h2>2. Data exploration and cleaning</h2>
<div class="box">The files were usable after {len(log)} fixes. The ones that would have changed answers: <b>250 duplicated customers</b>, <b>162 duplicated campaign contacts</b>,
<b>40 duplicated tickets</b>, contract end dates written three ways (2026-01-31, 31/01/2026 and 2026年1月31日; {K['mixed_dates']} rows), and
<b>{n(K['late_contacts'])} contacts whose 30-day tracking window had not finished</b> by 30 June, which would have made recent campaigns look worse.
Agent tags are unreliable: {n(K['missed'])} notes describe a customer about to leave but carry another tag.</div>
{table(log, ['file','step','rows_before','rows_after','note'], ['File','Fix','Rows before','Rows after','Detail'], {'rows_before': n, 'rows_after': n})}
<h3>Values we could not trust, and what we did</h3>
<ul>
<li><b>Latest values only.</b> Plan, contract end and tenure describe 1 July (or the last day), not the past. We never use contract end or plan to pick customers in the backtest except for the comparison rule, and tenure is recounted to each test date.</li>
<li><b>data_gb = 999</b> means unlimited; it is a flag, not 999 GB. 750 customers have no allowance recorded and are treated as unknown.</li>
<li><b>Vague campaign targets.</b> {K['vague']} past campaigns were briefed as "TBD", "everyone" or "not yet defined".</li>
<li><b>Checks that keep problems out:</b> every run re-counts rows at each step; any note sentence the AI has not read sends that note to a review queue as "unknown" instead of guessing.</li>
</ul></div>""")

    # ---------------- Part 3 ----------------
    seg_t = seg.copy()
    seg_t["plan_type"] = seg_t["plan_type"].map(PLAN_NAMES)
    seg_t = seg_t.set_index("plan_type").loc[list(PLAN_NAMES.values())].reset_index()
    parts.append(f"""
<div class="part"><h2>3. Analysis</h2>
<div class="box">Postpaid churn runs at about {pct(K['monthly_mean'], 1)} a month, as the fact sheet says. What customers tell support is by far the strongest early warning in the data:
customers whose notes show leave intent leave {K['sig_lift']:.0f}× more often than average. New customers (first six months) are the second warning. Price and rival offers are the main reason flagged customers give.
Past campaigns overstate their effect: once compared with the control group, roaming adds nothing and upsell returns less than half of what was reported.</div>
<h3>The customer base</h3>
{table(seg_t, ['plan_type','customers','customers_full_base','avg_bill','churn_6m'], ['Plan','Customers (sample)','Customers (full base)','Average bill','Left Jan–Jun 2026'], {'customers': n, 'customers_full_base': n, 'avg_bill': twd, 'churn_6m': pct})}
<p class="small">Family and business customers belong to other teams: {n(K['flag_fb'])} of them show leave intent in notes and are listed in flagged_family_business.csv for those teams, never targeted.</p>
{ch.img('2_monthly_churn')}
<h3>Early warning signs: tested on the past</h3>
<p>We pretended it was 1 March, 1 April and 1 May 2026, read only the notes written in the 90 days before that day, put every postpaid customer into one tier, and checked who left in the next 60 days.</p>
{ch.img('1_backtest')}
{table(R['bt_table'], ['rule','customers','left_next_60d','churn_rate','lift_vs_all','share_of_all_leavers'], ['Rule','Customers (3 tests)','Left in 60 days','Churn rate','× average','Share of all leavers'], {'customers': n, 'left_next_60d': n, 'churn_rate': pct, 'lift_vs_all': lambda v: f"{v:.1f}×", 'share_of_all_leavers': pct})}
<p class="small">The note tiers are small but sharp: about 1 in 2 leave. They catch {pct(K['sig_share'])} of all leavers, because most leavers never call support. R3 widens the net at lower precision.</p>
<h3>Why they leave</h3>
{ch.img('3_reasons')}
<h3>What customers say about rivals</h3>
<p>Notes name a rival {n(K['rival_sentences'])} times. Customers who name one left at {pct(K['rival_churn_min'], 0)}–{pct(K['rival_churn_max'], 0)} in January–June, against {pct(K['six_month_all'], 0)} for all customers.
More than a third are already acting (a quote, a store visit, a porting request), so speed matters more than the size of the offer.
Where the rival simply offers more for the money, Data boost (NT$300) is the cheapest approved answer.</p>
{table(R['rival_themes'], ['theme','sentences','share'], ['What the note says','Sentences','Share'], {'sentences': n, 'share': pct})}
{table(R['rivals'], ['rival','customers','mentions','churn_6m'], ['Rival named','Customers (sample)','Mentions','Left Jan–Jun 2026'], {'customers': n, 'mentions': n, 'churn_6m': pct})}
<h3>Room to earn more</h3>
<p>{n(K['up_aud'])} customers at full scale asked about more data, a bigger plan or roaming in a note in the last 90 days. That is a short, warm list; past broad blasts went to tens of thousands with no reason to buy.</p>
<h3>How acquisition performs</h3>
{ch.img('5_paid')}
{table(R['paid_by_ch'].sort_values('cost_per_line'), ['channel','campaigns','spend','new_lines','cost_per_line','margin_roas_6m','median_payback_months'], ['Channel','Campaigns','Spend','New lines','Cost per line','6-month margin return','Payback (months)'], {'campaigns': n, 'spend': twd, 'new_lines': n, 'cost_per_line': twd, 'margin_roas_6m': lambda v: f"{v:.2f}×", 'median_payback_months': lambda v: f"{v:.1f}"})}
<h3>Do the new customers stay?</h3>
<p>We followed the new lines each paid channel brought in (sample) to 1 July 2026. Between {pct(R['acq_q']['still_with_us'].min(), 0)} and {pct(R['acq_q']['still_with_us'].max(), 0)} were still customers, so no channel brings in clearly worse customers.
Counting only customers who stayed, Google Search is still the cheapest. The sample is small (about 90–160 lines per channel), so the gaps between channels are indicative only.</p>
{table(R['acq_q'], ['channel','new_lines_sample','still_with_us','cost_per_line','cost_per_line_still_with_us'], ['Channel','New lines (sample)','Still a customer 1 Jul','Cost per new line','Cost per customer who stayed'], {'new_lines_sample': n, 'still_with_us': pct, 'cost_per_line': twd, 'cost_per_line_still_with_us': twd})}
</div>""")

    # ---------------- Part 4 ----------------
    val = R["val"]
    parts.append(f"""
<div class="part"><h2>4. Interpretation: what it is worth</h2>
<div class="box">At Formosa's full scale the retention play is worth {twd(K['churn_low'])}–{twd(K['churn_high'])} net margin a year, the earn-more play {twd(K['up_low'])}–{twd(K['up_high'])},
and the acquisition shift {twd(K['acq_low'])}–{twd(K['acq_high'])}. The number that would change our mind is the save rate: nobody knows it yet, which is why the pilot measures it first.</div>
<h3>1 · Churn</h3>
{table(val[val.objective=='1 churn'], ['scenario','save_rate','flagged_per_month_full_base','customers_kept_per_year','credited_revenue_twd','offer_cost_twd','net_margin_twd'], ['Case','Save rate','Flagged / month','Kept / year','Credited revenue','Offer cost','Net margin'], {'save_rate': pct, 'flagged_per_month_full_base': n, 'customers_kept_per_year': n, 'credited_revenue_twd': twd, 'offer_cost_twd': twd, 'net_margin_twd': twd})}
<h3>2 · Earn more</h3>
{table(val[val.objective=='2 earn more'], ['scenario','take_rate','audience_per_month_full_base','buyers_per_year','credited_revenue_twd','offer_cost_twd','net_margin_twd'], ['Case','Take rate','Audience / month','Buyers / year','Credited revenue','Offer cost','Net margin'], {'take_rate': pct, 'audience_per_month_full_base': n, 'buyers_per_year': n, 'credited_revenue_twd': twd, 'offer_cost_twd': twd, 'net_margin_twd': twd})}
<p class="small">Plus {twd(K['roam_spend_year'])} a year of roaming campaign spend that caused no extra sales. The earn-more range in Part 1 adds this saving.</p>
<h3>3 · Win new</h3>
{table(R['acq'], ['annual_paid_spend_twd','lines_if_all_best','lines_if_all_worst','extra_lines_from_shift','extra_6m_margin_from_shift_twd'], ['Paid spend / year','Lines at Google Search cost','Lines at Retail cost','Extra lines','Extra 6-month margin'], {'annual_paid_spend_twd': twd, 'lines_if_all_best': n, 'lines_if_all_worst': n, 'extra_lines_from_shift': n, 'extra_6m_margin_from_shift_twd': twd})}
<p class="small">Upper bound: not all spend can move to one channel. Counted on Finance's six-month rule; customers who stay longer are worth more.</p>
<h3>The assumptions that matter most</h3>
{table(R['assump'], ['what','value','why'], ['Assumption','Value','Why'])}
<p><b>What would change our mind:</b> if the pilot's held-back group leaves at the same rate as the contacted group by day 45, retention aimed at notes does not work at these offers, and the money should go to fixing the billing and network causes instead.</p>
</div>""")

    # ---------------- Part 5 ----------------
    ca = R["copy_audit"]
    bad = ca[~ca["compliant"]]
    parts.append(f"""
<div class="part"><h2>5. Campaign engine</h2>
<div class="box">Audiences are picked by written rules in code (signals.py), refreshed whenever a new notes file arrives. Every message is drafted from approved offers only,
gets the legal footer added by code, and passes an automatic compliance check. {K['bad_copy']} of {K['all_copy']} past campaigns would have failed that check.</div>
<h3>The rules</h3>
<table><tr><th>Tier</th><th>Rule (only notes from the 90 days before the send)</th><th>What we send</th></tr>
<tr><td>R1 leaving now</td><td>Note says cancelling, porting or signing elsewhere</td><td>Outbound call, highest bill first; matched retention offer</td></tr>
<tr><td>R2 wavering</td><td>Note says comparing, unsure about renewal, or will leave if X</td><td>Call if capacity, else LINE/SMS; matched retention offer</td></tr>
<tr><td>R3 new</td><td>Joined less than six months ago</td><td>Welcome check-in, no paid offer</td></tr>
<tr><td>R4 unhappy, quiet</td><td>Billing or network complaint, no leave talk</td><td>Fix it: billing care or network care, no paid offer</td></tr>
<tr><td>Excluded</td><td>Line closing for a reason no offer changes; family and business plans</td><td>Nothing</td></tr></table>
<p class="small">Offer match: device problem on a premium plan → Device credit; price, rival or renewal doubt within four months of contract end → Renewal discount; price or data need and not unlimited → Data boost; billing error → fix it (not an offer); network → care call.</p>
<h3>New campaigns</h3>
{table(camp, ['campaign','channel','customers_at_full_scale','timing','headline','compliant'], ['Audience | offer','Channel','Customers (full base)','Timing','Headline','Compliant'], {'customers_at_full_scale': n})}
<h3>Example message (as sent, {ex['channel']})</h3>
<div class="box" style="white-space:pre-wrap">{ex['headline']}\n\n{ex['body']}\n\n[ {ex['cta']} ]</div>
<h3>Marketer tool</h3>
<p>A marketer types a brief and gets the audience size, the copy and a compliance result: <code>py make_campaign.py "renewal discount by SMS for customers who mention a rival"</code></p>
<h3>Audit of past campaigns</h3>
{table(bad.groupby('problems').size().reset_index(name='campaigns').sort_values('campaigns', ascending=False), ['problems','campaigns'], ['Problem','Campaigns'], {'campaigns': n})}
<h3>QA checklist before any send</h3>
<ul><li>Offer is on the approved list for this objective; no free months, cash or bill waivers.</li>
<li>Legal footer verbatim as the closing paragraph (short terms line on SMS and LINE only).</li>
<li>No banned claim or close paraphrase ("never drops", "cheapest", "fastest network" unqualified…).</li>
<li>Audience excludes family, business, prepaid (for retention) and "no save" lines; 20% held back at random.</li>
<li>Outbound calls within 2,000 a month; prices in NT$; copy in English.</li></ul>
</div>""")

    # ---------------- Part 6 ----------------
    parts.append(f"""
<div class="part"><h2>6. Measurement and the pilot</h2>
<div class="box">Past owned campaigns earned less than reported: against the group held back on purpose, roaming caused nothing and upsell caused {pct(roas.loc[roas.objective=='upsell','share_of_reported_that_is_real'].iloc[0], 0)} of what it claimed.
The pilot is built so this mistake cannot happen again: 20% held back at random from day one, success measured as the gap in churn between the two groups.</div>
{ch.img('4_roas')}
{table(roas, ['objective','campaigns','spend','reported_revenue','reported_roas','incremental_roas','incremental_margin_roas'], ['Objective','Campaigns','Spend','Reported revenue','Reported ROAS','Real ROAS','Real margin return'], {'campaigns': n, 'spend': twd, 'reported_revenue': twd, 'reported_roas': lambda v: f"{v:.2f}×", 'incremental_roas': lambda v: f"{v:.2f}×", 'incremental_margin_roas': lambda v: f"{v:.2f}×"})}
<h3>The 90-day pilot</h3>
<table><tr><th>Item</th><th>Plan</th></tr>
<tr><td>Who</td><td>Postpaid basic and premium customers in tiers R1 and R2, refreshed weekly from new notes. About {n(mid['flagged_full_base_90d'])} customers over 90 days at full scale.</td></tr>
<tr><td>Held back</td><td>{n(mid['held_back_20pct'])} (one in five, at random). They get nothing.</td></tr>
<tr><td>Channels</td><td>{n(mid['outbound_calls'])} outbound calls (2,000 a month, highest bill first); everyone else LINE and SMS.</td></tr>
<tr><td>Offers</td><td>Renewal discount, Data boost or Device credit by reason; billing and network problems fixed, not discounted.</td></tr>
<tr><td>Cost</td><td>{twd(mid['total_cost'])} expected (messages {twd(mid['messaging_cost'])}, offers {twd(mid['offer_cost'])}, creative {twd(mid['creative_cost'])}); at most {twd(K['pilot_cost_high'])}. Hard cap NT$2.5M: offers stop when the cap is reached.</td></tr>
<tr><td>Expected</td><td>Keep {n(p.iloc[0]['customers_kept'])}–{n(p.iloc[2]['customers_kept'])} customers; credited revenue {twd(p.iloc[0]['credited_revenue'])}–{twd(p.iloc[2]['credited_revenue'])}.</td></tr>
<tr><td>Break-even</td><td>Keep {pct(K['breakeven'])} of would-be leavers: churn in the contacted group {pp(K['breakeven_gap'], 2)} percentage points below the held-back group.</td></tr>
<tr><td>What it can prove</td><td>With {n(mid['held_back_20pct'])} held back, chance alone moves the gap by up to ±{pp(K['noise_95'], 1)} percentage points. The pilot can reliably prove a gap of {pp(K['mde_gap'], 1)} percentage points or more, i.e. a save rate of {pct(K['mde_save'])} or more (95% confidence, 80% power). A save rate between {pct(K['breakeven'])} and {pct(K['mde_save'])} would still make money but cannot be proven in 90 days.</td></tr>
<tr><td>Day 45 check</td><td>Stop paid offers (keep care calls) if the contacted group's churn is not below the held-back group's. At day 45 noise is about ±{pp(K['noise_95_day45'], 1)} percentage points, so this checks the direction only, not the size.</td></tr>
<tr><td>Day 90 verdict</td><td>Gap of {pp(K['mde_gap'], 1)} percentage points or more: proven, scale to the full base. Gap above zero but smaller: profitable but unproven, extend 90 days to add customers. Gap at or below zero: stop.</td></tr>
<tr><td>Measure</td><td>Churn within 60 days of contact, contacted vs held back; offer cost per kept customer; take-up by offer and channel.</td></tr></table>
{ch.img('6_pilot')}
<h3>Risks</h3>
<ul><li><b>Save rate unknown.</b> Formosa has never run reason-matched retention. The held-back group measures it, but only effects of {pct(K['mde_save'], 0)} or more can be proven in 90 days.</li>
<li><b>Offers taken by people who would have stayed.</b> Costed at 1.5 takers per customer kept.</li>
<li><b>Coverage.</b> Notes catch about {pct(K['sig_share'], 0)} of leavers. Next phase: the R3 welcome check-in, and asking agents to log every contact.</li>
<li><b>Call centre capacity.</b> 2,000 calls a month; the rest go to LINE and SMS, so the result is measured by channel.</li></ul>
<h3>How we used AI</h3>
<p class="small">Claude Opus 5.5 read the 403 distinct sentences that make up all 8,000 notes once and labelled each with a topic and a leave-intent level; it also wrote the message templates. The code only looks these up, so the results rerun offline. See 3_ai_outputs/how_we_did_it.md.</p>
</div>""")

    return f"<!DOCTYPE html><html><head><meta charset='utf-8'><title>Formosa Telecom proposal</title><style>{CSS}</style></head><body>{''.join(parts)}</body></html>"


def slides_html(R, K, ch):
    mid = R["pilot"][R["pilot"]["scenario"] == "mid"].iloc[0]
    css = """@page { size: 297mm 167mm; margin: 0; } body { margin: 0; font-family: 'Segoe UI', Helvetica, Arial, sans-serif; color: #1f2933; }
.s { width: 297mm; height: 167mm; box-sizing: border-box; padding: 12mm 16mm; page-break-after: always; position: relative; }
.s h1 { font-size: 30pt; margin: 0 0 6mm; color: #0b6e99; } .n { position: absolute; top: 10mm; right: 16mm; color: #7b8794; font-size: 10pt; }
.big { font-size: 19pt; line-height: 1.45; } .cols { display: flex; gap: 10mm; } .cols > div { flex: 1; }
.k b { display: block; font-size: 32pt; color: #0b6e99; } .k { font-size: 14pt; margin-bottom: 7mm; } img { width: 100%; max-height: 120mm; object-fit: contain; }
ul { font-size: 17pt; line-height: 1.6; padding-left: 7mm; } .ask { background: #fff8e6; border-left: 6px solid #d69e2e; padding: 6mm; font-size: 20pt; margin: 6mm 0; }"""
    s = []
    s.append(f"""<div class="s"><div class="n">1 / 6</div><h1>The problem and our approach</h1><div class="cols"><div class="big">
Vivian needs to know who is leaving, who will buy more, and which acquisition pays back, in NT$, with proof.<br><br>
<b>Our approach:</b> read what customers already tell support (8,000 notes nobody reads), test every rule on past dates, and measure every campaign against a group held back on purpose.</div>
<div><ul><li>Cleaned 5 files, logged every fix</li><li>AI read every note sentence once</li><li>Rules tested on 1 Mar, 1 Apr, 1 May</li><li>Past campaigns vs control group</li><li>One command rebuilds it all</li></ul></div></div></div>""")
    s.append(f"""<div class="s"><div class="n">2 / 6</div><h1>What the data shows</h1><div class="cols"><div style="flex:1.6">{ch.img('1_backtest')}</div>
<div><div class="k"><b>{pct(K['sig_rate'], 0)}</b>of customers who tell support they may leave are gone in 60 days</div>
<div class="k"><b>{pct(K['ce_rate'])}</b>for today's contract-end rule (average {pct(K['base_rate'])})</div>
<div class="k"><b>0</b>extra sales from roaming blasts vs control</div></div></div></div>""")
    s.append(f"""<div class="s"><div class="n">3 / 6</div><h1>Who to target, and why</h1><div class="cols"><div>{ch.img('3_reasons')}</div>
<div><ul><li><b>R1 leaving now</b>: porting, cancelling, signing elsewhere → call first</li><li><b>R2 wavering</b>: comparing, unsure on renewal → call or LINE/SMS</li>
<li><b>R3 new</b> (&lt; 6 months, {K['r3_lift']:.1f}× churn) → welcome check-in</li><li><b>R4 unhappy</b> → fix the bill or the network</li>
<li><b>Earn more</b>: only customers who asked ({n(K['up_aud'])} at full scale)</li><li>Rival named in a note → {pct(K['rival_churn_min'], 0)}–{pct(K['rival_churn_max'], 0)} left; a third already getting quotes or porting</li></ul></div></div></div>""")
    s.append(f"""<div class="s"><div class="n">4 / 6</div><h1>The campaigns</h1><div class="cols"><div><ul>
<li>Reason → approved offer: rival/price → Renewal discount or Data boost; device → Device credit</li><li>Billing and network: fix it, no discount</li>
<li>Stop roaming blasts ({twd(K['roam_spend_year'])}/yr, no effect)</li><li>Every message: legal footer + automatic compliance check</li>
<li>{K['bad_copy']} of {K['all_copy']} past campaigns would have failed</li></ul></div>
<div class="k"><b>{twd(K['churn_low'])}–{twd(K['churn_high'])}</b>a year, retention<br><br><b>{twd(K['up_low'])}–{twd(K['up_high'])}</b>a year, earn more<br><br><b>{twd(K['acq_low'])}–{twd(K['acq_high'])}</b>a year, acquisition shift</div></div></div>""")
    s.append(f"""<div class="s"><div class="n">5 / 6</div><h1>The 90-day pilot</h1><div class="cols"><div style="flex:1.4">{ch.img('6_pilot')}</div>
<div><div class="k"><b>{twd(mid['total_cost'])}</b>expected cost, max {twd(K['pilot_cost_high'])}</div><div class="k"><b>{pct(K['breakeven'])}</b>of would-be leavers kept = break-even</div>
<div class="k"><b>{pct(K['mde_save'], 0)}+</b>save rate the pilot can prove in 90 days (20% held back)</div></div></div></div>""")
    s.append(f"""<div class="s"><div class="n">6 / 6</div><h1>Our recommendation</h1><div class="ask">Approve the 90-day retention pilot: {twd(mid['total_cost'])} expected, inside NT$2.5M,
for postpaid customers who tell support they may leave.</div><ul><li>Week 1: connect the weekly notes file; first call list to the call centre</li>
<li>Day 45: stop paid offers if contacted customers leave as often as held-back ones</li><li>Day 90: proven if the gap is {pp(K['mde_gap'], 1)} percentage points or more</li><li>Then: scale retention to the full base; redirect roaming and broad upsell spend</li></ul></div>""")
    return f"<!DOCTYPE html><html><head><meta charset='utf-8'><title>Finals slides</title><style>{css}</style></head><body>{''.join(s)}</body></html>"


# ---------- finals talk script ----------

def talk_script(R, K):
    mid = R["pilot"][R["pilot"]["scenario"] == "mid"].iloc[0]
    worst = mid["offer_cost"] / 2 + mid["messaging_cost"] + mid["creative_cost"]
    return f"""# Finals talk script (8 minutes) and Q&A

Two presenters: **A** (story, slides 1, 2, 6) and **B** (method, slides 3, 4, 5). About 80 seconds a slide.
Every number below comes from the pipeline output; rerun `python run.py` and this file updates.

## Slide 1 · The problem and our approach (A)
"In the three months before 1 April, {K['r1_apr_n']} postpaid customers in our 1-in-130 sample, about {n(K['r1_apr_n'] * 130)} across Formosa,
told a support agent they were cancelling, porting their number or signing with a rival. {K['r1_apr_left']} of those {K['r1_apr_n']} were gone within 60 days.
Nobody at Formosa reads these notes. Vivian asked who is leaving, why, and what keeps them. Her customers are already telling her."

## Slide 2 · What the data shows (A)
- Customers whose notes show leave intent: {pct(K['sig_rate'])} leave within 60 days. Average: {pct(K['base_rate'])}.
- Today's renewal push (contracts ending within 4 months): {pct(K['ce_rate'])}. Barely better than picking at random.
- We tested this on three past dates using only notes written before each date, so it is not hindsight.
- Roaming blasts: no extra sales against the held-back group.

## Slide 3 · Who to target, and why (B)
- R1 leaving now and R2 wavering: call the highest bills first, then LINE and SMS.
- The reason decides the offer: price or a rival → Renewal discount or Data boost; phone problem → Device credit; bill error or bad signal → fix it, no discount.
- Rivals are named {n(K['rival_sentences'])} times; those customers left at {pct(K['rival_churn_min'], 0)}–{pct(K['rival_churn_max'], 0)}. A third were already getting quotes or porting, so speed beats discount size.

## Slide 4 · The campaigns (B)
- Every message: approved offer, legal footer added by code, automatic compliance check. {K['bad_copy']} of {K['all_copy']} past campaigns would have failed it.
- Worth a year at full scale: retention {twd(K['churn_low'])}–{twd(K['churn_high'])}; earn more {twd(K['up_low'])}–{twd(K['up_high'])}; acquisition shift {twd(K['acq_low'])}–{twd(K['acq_high'])}.
- Live demo (30 seconds): open `4_outputs/campaign_builder.html`, pick Renewal discount + SMS + "Price or a rival", show the audience size and the PASS.

## Slide 5 · The 90-day pilot (B)
- About {n(mid['flagged_full_base_90d'])} flagged postpaid customers; one in five held back at random.
- Expected cost {twd(mid['total_cost'])}, at most {twd(K['pilot_cost_high'])}, inside NT$2.5M.
- Break-even: keep {pct(K['breakeven'])} of would-be leavers. The pilot can prove {pct(K['mde_save'])} or more.
- Day 45: if contacted customers leave as often as held-back ones, paid offers stop.

## Slide 6 · Our recommendation (A)
"Approve the pilot. Week one, we connect the weekly notes file and send the first call list.
Day 45, we check direction. Day 90, if the gap is {pp(K['mde_gap'], 1)} percentage points or more, scale it to the whole base and move roaming and broad upsell money into it."

---

## Hard questions, short answers

**What if the save rate is zero?**
Then we find out cheaply. Offers cost money only when taken; the day-45 check stops them. Worst case is roughly half the expected offer cost plus messages and creative, about {twd(worst)}, and Formosa learns that note-based retention does not work at these offers.

**Why rules and not a machine-learning model?**
Customers the rules flag leave {K['sig_lift']:.0f}× more often than average on dates the rules never saw, a marketer can read and change them, and every customer comes with a reason that decides the offer. A model on 50,000 rows whose plan and contract fields show only the latest values would risk learning the future.

**How do you know the AI read the notes correctly?**
It labelled 403 distinct sentences, not 8,000 notes, so every label can be checked by a person. {K['hand_check_text']} The outcome agrees: flagged customers left at {pct(K['sig_rate'])}. Any sentence it has never seen goes to a review queue as "unknown".

**Notes catch only {pct(K['sig_share'], 0)} of leavers. Isn't that small?**
It is the precise part: about 1 in 2 flagged customers leave. Most leavers never call support. Next: the welcome check-in for new customers ({K['r3_lift']:.1f}× average churn) and asking agents to log every contact.

**Why multiply by 130?**
The fact sheet says the customer files are a fair 1-in-130 sample of 6.5 million customers.

**Can you trust the past-campaign audit?**
Since July 2025 one in five chosen customers was held back at random. We compare only those campaigns, and only contacts whose 30-day tracking window had finished.

**Isn't Renewal discount expensive?**
It goes only to customers within four months of contract end who raised price or renewal. Customers who only want more data get Data boost at NT$300.

**Why can't the pilot prove the break-even?**
With {n(mid['held_back_20pct'])} held back, chance alone moves the result by about ±{pp(K['noise_95'], 1)} percentage points, more than the {pct(K['breakeven_gap'], 2)}-point break-even gap. We say so up front and give a rule for each outcome instead of over-promising.
"""


# ---------- clickable campaign builder ----------

def campaign_builder(R, templates_path):
    """One self-contained page: pick offer, channel, tier, reason → audience size, message, compliance."""
    import json
    from compliance import FOOTER, SHORT_FOOTER, BANNED
    m = R["members"].copy()
    m["reason"] = m["reasons"].fillna("").map(lambda r: next((x for x in ["contract", "price", "device", "billing", "network", "data_need"] if x in r.split(";")), "other"))
    counts = m.groupby(["offer", "tier", "reason"]).size().reset_index(name="n").to_dict("records")
    with open(templates_path, encoding="utf-8") as f:
        tpl = json.load(f)
    data = {"counts": counts, "templates": {k: v for k, v in tpl.items() if not k.startswith("_")},
            "footer": FOOTER, "short": SHORT_FOOTER, "banned": BANNED,
            "share": {"SMS": 0.74, "LINE": 0.41, "Email": 0.33, "App Push": 0.27, "Outbound Call": 0.92}}
    return """<!DOCTYPE html><html><head><meta charset="utf-8"><title>Formosa campaign builder</title>
<style>body{font-family:'Segoe UI',Arial,sans-serif;color:#1f2933;max-width:900px;margin:24px auto;padding:0 16px}
h1{color:#0b6e99;font-size:22px}label{display:block;font-size:13px;color:#52606d;margin-top:10px}
select{width:100%;padding:8px;font-size:14px;border:1px solid #cbd2d9;border-radius:6px}
.grid{display:grid;grid-template-columns:1fr 1fr;gap:12px}.card{border:1px solid #e4e7eb;border-radius:8px;padding:14px;margin-top:16px}
.big{font-size:26px;color:#0b6e99;font-weight:600}.msg{white-space:pre-wrap;background:#f5f7fa;padding:12px;border-radius:6px;font-size:14px}
.pass{color:#2f855a;font-weight:600}.fail{color:#c53030;font-weight:600}button{padding:8px 14px;border:0;border-radius:6px;background:#0b6e99;color:#fff;cursor:pointer}</style></head>
<body><h1>Formosa campaign builder</h1><p>Pick an approved offer and a channel. You get the audience, the finished message with the legal footer, and a brand check. Works offline.</p>
<div class="grid"><div><label>Offer</label><select id="offer"></select></div><div><label>Channel</label><select id="channel"></select></div>
<div><label>Who</label><select id="tier"></select></div><div><label>Reason in their note</label><select id="reason"></select></div></div>
<div class="card"><div class="big" id="size"></div><div id="sizeNote"></div></div>
<div class="card"><b id="head"></b><div class="msg" id="body"></div><p><b>Call to action:</b> <span id="cta"></span></p>
<p id="check"></p><button onclick="navigator.clipboard.writeText(document.getElementById('head').innerText+'\\n\\n'+document.getElementById('body').innerText)">Copy message</button></div>
<script>const D=""" + json.dumps(data) + """;
const TIERS={any:"Anyone with this offer",R1_leaving_now:"R1 · said they are leaving",R2_wavering:"R2 · considering leaving",R3_new_and_fragile:"R3 · joined < 6 months ago",R4_unhappy_quiet:"R4 · complaint, no leave talk"};
const REASONS={any:"Any reason",contract:"Unsure about renewing",price:"Price or a rival",device:"Phone problem",billing:"Bill looks wrong",network:"Coverage or speed",data_need:"Runs out of data",other:"Other"};
const RT={contract:"they are unsure about renewing",price:"they felt the plan is expensive or that a rival offers more",billing:"their bill looked wrong",network:"signal problems",device:"a problem with their phone",data_need:"running out of data"};
const $=id=>document.getElementById(id);
function fill(id,obj){$(id).innerHTML=Object.entries(obj).map(([k,v])=>`<option value="${k}">${v}</option>`).join("")}
fill("offer",Object.fromEntries(Object.keys(D.templates).map(k=>[k,k])));fill("channel",Object.fromEntries(Object.keys(D.share).map(k=>[k,k])));fill("tier",TIERS);fill("reason",REASONS);
function check(h,b,c,ch){const p=[];const t=[h,b,c].join(" ");const end=b.trim();
 if(!(end.endsWith(D.footer)||((ch==="SMS"||ch==="LINE")&&end.endsWith(D.short))))p.push("legal footer missing or altered");
 for(const[k,re]of Object.entries(D.banned)){if(new RegExp(re,"i").test(t)&&!(k.startsWith("fastest")&&/independent|20\\d\\d/i.test(t)))p.push("banned claim: "+k)}
 if((t.match(/!/g)||[]).length>=2)p.push("stacked exclamation marks");return p}
function update(){const o=$("offer").value,ch=$("channel").value,ti=$("tier").value,re=$("reason").value,T=D.templates[o];
 const n=D.counts.filter(r=>r.offer===o&&(ti==="any"||r.tier===ti)&&(re==="any"||r.reason===re)).reduce((a,r)=>a+r.n,0);
 $("size").innerText=(n*130).toLocaleString()+" customers";
 $("sizeNote").innerText=`${n} in the sample × 130. About ${Math.round(n*130*D.share[ch]).toLocaleString()} can be reached on ${ch} (${Math.round(D.share[ch]*100)}% have agreed).`+(n?"":" Nobody matches: widen Who or Reason.");
 let body;if(ch==="Outbound Call"){body=(T.call_script||T.body).replace("{reason_text}",RT[re]||"their recent contact with us")}
 else{const f=(ch==="SMS"||ch==="LINE")?D.short:D.footer;body=(ch==="SMS"?"Formosa: "+T.body+" "+T.cta+".":T.body)+"\\n\\n"+f}
 $("head").innerText=T.headline;$("body").innerText=body;$("cta").innerText=T.cta;
 const p=ch==="Outbound Call"?[]:check(T.headline,body,T.cta,ch);
 $("check").innerHTML=ch==="Outbound Call"?'<span class="pass">Call script</span> · agents read the terms on request':(p.length?'<span class="fail">FAIL</span> · '+p.join("; "):'<span class="pass">PASS</span> · approved offer, legal footer verbatim, no banned claims')}
["offer","channel","tier","reason"].forEach(i=>$(i).onchange=update);update();</script></body></html>"""
